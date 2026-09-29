"""Chile (IPC base 2023): extrai a cesta (produtos) e o gabarito COICOP a partir da Planilha2 (compatibilização manual)."""
import openpyxl, pandas as pd
wb = openpyxl.load_workbook("data_chile.xlsx", read_only=True, data_only=True)
# gabarito manual: (código Chile, glosa Chile) -> COICOP  (linhas com vários pares: código,glosa a cada 3 colunas)
gold = {}
for r in wb["Planilha2"].iter_rows(values_only=True):
    if not r[1]: continue
    cop = str(r[1]).split()[0]
    for i in range(2, len(r) - 1, 3):
        if r[i] is not None and r[i + 1]: gold[(str(r[i]), r[i + 1].strip())] = cop
# cesta 2024/01: reconstrói o código concatenado dos ancestrais de cada produto
rows, anc = [], [None] * 5
for r in wb["IPC 2023=100"].iter_rows(min_row=5, values_only=True):
    if r[0] != 2024 or r[1] != 1: continue
    lv = [x for x in r[2:7]]; glosa = r[7]
    n = sum(x is not None for x in lv)              # 0=IPC geral, 1=divisão, ..., 5=produto
    if n == 0: continue
    lv = lv[:n]; anc[:n] = [str(x) for x in lv]; 
    rows.append(dict(nivel=n, codigo="".join(anc[:n]), glosa=glosa.strip(), ancestrais=list(anc[:n]), peso=r[8]))
df = pd.DataFrame(rows)
names = {(x.nivel, x.codigo): x.glosa for x in df.itertuples()}
def rotulo(x):   # classe COICOP pelo ancestral mais específico que consta do gabarito
    for k in range(x.nivel, 0, -1):
        cod = "".join(x.ancestrais[:k]); g = names[(k, cod)]
        if (cod, g) in gold: return gold[(cod, g)], k
    return "SEM_CORRESPONDENCIA", 0
prod = df[df.nivel == 5].copy()
prod[["coicop", "nivel_gabarito"]] = [rotulo(x) for x in prod.itertuples()]
prod["subitem"] = prod.glosa
prod["item"] = [names[(4, "".join(a[:4]))] for a in prod.ancestrais]
prod["grupo"] = [names[(1, a[0])] for a in prod.ancestrais]
prod["item_cod"] = ["".join(a[:4]) for a in prod.ancestrais]
prod = prod[["codigo", "grupo", "item_cod", "item", "subitem", "peso", "coicop", "nivel_gabarito"]]
prod.to_csv("chile_produtos_rotulados.csv", index=False)
print("produtos:", len(prod), "| sem correspondência:", (prod.coicop == "SEM_CORRESPONDENCIA").sum(),
      "| classes COICOP:", prod.coicop.nunique(), "| peso coberto: %.1f%%" % (100 * prod.peso[prod.coicop != "SEM_CORRESPONDENCIA"].sum() / prod.peso.sum()))
print(prod[prod.coicop == "SEM_CORRESPONDENCIA"][["codigo", "item", "subitem"]].head(15).to_string())
