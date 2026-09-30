"""Corrige o de-para SNIPC -> COICOP 1999 do Brasil (mapeamentos/BRA.csv) para o conteúdo OFICIAL de cada classe de 1999 (notas da UNSD,
aba 'COICOP 1999' de coicop_ml/fontes/coicop2018_1999_correspondencia.xlsx). Gera dados/Correcoes_BRA_COICOP1999.xlsx (antes/depois).
Só entram casos com evidência no texto oficial; os demais subitens NÃO foram revisados linha a linha."""
import sys, pandas as pd
from pathlib import Path
R = Path(__file__).resolve().parents[1]; sys.path.insert(0, str(R / "src"))
from harmonizacao import armazem as A
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
m = pd.read_csv(R / "mapeamentos/BRA.csv", dtype=str)
un = pd.read_excel(R / "../coicop_ml/fontes/coicop2018_1999_correspondencia.xlsx", sheet_name="COICOP 1999", dtype=str)
tit = dict(zip(un.Code, un.Description.str.replace(r"\s*\((ND|SD|D|S)\)$", "", regex=True)))
# (códigos SNIPC, classe antes esperada, classe depois, motivo com base no texto oficial)
C = [
 ("6301001 6301002 6301004 6301006 6301007 6301010 6301011 6301014 6301015 6301016 6301017 6301020", "12.1.2", "12.1.3",
  "12.1.2 oficial é só 'Electric appliances for personal care'. Produtos de higiene e beleza, perfume, desodorante, absorvente, papel higiênico e fralda ('babies' napkins') são 12.1.3."),
 ("2104041", "05.6.1", "12.1.3", "12.1.3 lista expressamente 'paper towels' (papel toalha)."),
 ("8101004 8101008", "10.1.0", "10.2.0", "10.1.0 oficial = níveis 0 e 1 da ISCED (pré-escola e fundamental). 10.2.0 = ensino médio e 'out-of-school secondary education for adults' (EJA)."),
 ("8101005 8101006", "10.1.0", "10.4.0", "10.4.0 = ensino superior (ISCED 5 e 6), onde entram graduação e pós-graduação."),
 ("8101045 8104001 8104003 8104004", "10.1.0", "10.5.0", "10.5.0 = programas para adultos sem pré-requisito, em especial formação profissional e desenvolvimento cultural (curso técnico, preparatório, idiomas, informática)."),
 ("7201019", "07.1.2", "07.1.3", "07.1.2 oficial = motocicletas; 07.1.3 = 'Bicycles and tricycles of all types'."),
 ("7201020", "09.3.4/5", "09.3.4", "09.3.4 = 'Pets, pet foods...'; código '09.3.4/5' não existe na lista oficial."),
 ("7201015 7201256", "09.3.4/5", "09.3.5", "09.3.5 = serviços veterinários e de higiene/tosa para animais."),
 ("7201063", "09.3.1", "09.4.3", "09.4.3 = 'Games of chance' (loterias, apostas, cassinos)."),
 ("9101002 9101008 9101018 9101116", "08.2.0", "08.3.0", "08.2.0 oficial = aparelhos (compra e reparo). 08.3.0 = serviços: assinatura, chamadas, internet e combos."),
 ("3301006", "05.3.3", "09.1.5", "05.3.3 é reparo de eletrodomésticos; reparo de televisor é 09.1.5 (equipamento audiovisual)."),
 ("7101001", "12.1.1", "03.1.4", "03.1.4 inclui 'darning, mending, repair and altering of garments' (costureira)."),
 ("1110009 1110010", "01.1.4", "01.1.2", "Frango é carne de ave (01.1.2). 01.1.4 é leite, queijo e ovos."),
 ("1101051 1101052 1101053 1101073", "01.1.1", "01.1.7", "01.1.7 inclui hortaliças secas e leguminosas; feijão não é cereal."),
 ("1111031", "01.1.4", "01.1.5", "01.1.5 = 'Butter and butter products'; 01.1.4 exclui manteiga."),
 ("1109023", "01.1.2", "01.1.3", "Bacalhau: 01.1.3 inclui 'dried, smoked or salted fish'."),
 ("1116010", "01.1.9", "01.1.7", "01.1.7 cita alho (garlic) entre as hortaliças."),
 ("1104052", "01.1.8", "01.2.1", "01.2.1 inclui 'cocoa... and chocolate-based powder'; 01.1.8 fica com chocolate em barra. (Em 2018 o mesmo produto vai para 01.1.8.)"),
]
peso = A.ler("peso_nacional"); peso = peso[(peso.pais == "BRA") & (peso.data == peso.data.max())].set_index("cod_local").peso
nome = A.ler("item_nacional").set_index("cod_local").nome
rows = []; novo = m.set_index("cod_local")
for cods, antes, depois, motivo in C:
    for c in cods.split():
        ant = novo.at[c, "coicop"]; assert ant == antes, f"{c}: esperado {antes}, encontrado {ant}"
        novo.at[c, "coicop"] = depois; novo.at[c, "metodo"] = "corrigido_texto_oficial_1999"; novo.at[c, "revisado"] = "False"
        rows.append({"Código SNIPC": c, "Subitem": nome[c], "Peso ago/2026 (%)": round(peso[c], 4), "COICOP 1999 antes": antes, "COICOP 1999 depois": depois,
                     "Título oficial (depois)": tit.get(depois, ""), "Motivo (texto oficial de 1999)": motivo})
tab = pd.DataFrame(rows).sort_values("Peso ago/2026 (%)", ascending=False)
assert set(novo.coicop) <= set(tit), sorted(set(novo.coicop) - set(tit))            # todos os códigos agora existem na lista oficial
novo.reset_index().to_csv(R / "mapeamentos/BRA.csv", index=False)
# impacto por classe (peso)
w = m.assign(peso=m.cod_local.map(peso)); antes = w.groupby("coicop").peso.sum(); depois = novo.assign(peso=novo.index.map(peso)).groupby("coicop").peso.sum()
imp = pd.DataFrame({"antes": antes, "depois": depois}).fillna(0); imp["variação"] = (imp.depois - imp.antes).round(4); imp = imp[imp["variação"].abs() > 1e-6].reset_index(names="COICOP 1999")
imp["Título oficial"] = imp["COICOP 1999"].map(tit); imp[["antes", "depois"]] = imp[["antes", "depois"]].round(4)
imp = imp.rename(columns={"antes": "Peso antes (%)", "depois": "Peso depois (%)", "variação": "Variação (p.p.)"}).sort_values("Variação (p.p.)", key=abs, ascending=False)
leia = pd.DataFrame({"Leia-me": [
 "Correções ao de-para SNIPC -> COICOP 1999 do Brasil, com base nas notas explicativas oficiais da UNSD para a COICOP 1999.",
 f"{len(tab)} subitens alterados, {tab['Peso ago/2026 (%)'].sum():.2f}% do peso do IPCA. O total do IPCA não muda: muda só a distribuição entre classes.",
 "O arquivo mapeamentos/BRA.csv foi atualizado (método 'corrigido_texto_oficial_1999', revisado = False). A versão anterior está no histórico do Git.",
 "LIMITE: só foram corrigidos casos com evidência no texto oficial. Os outros subitens NÃO foram revisados um a um contra as notas de 1999.",
 "Dúvida deixada como está: Leite de coco (1116001). A nota oficial de 01.1.4 cita 'dairy products not based on milk such as soya milk', mas não cita coco; peso 0,002%.",
 "Em 2018 alguns destes produtos caem em outra classe (ex.: chocolate em pó vai para 01.1.8); a classificação em 2018 fica na planilha COICOP2018_SNIPC_Brasil_e_Mercosul.xlsx."]})
out = R / "dados/Correcoes_BRA_COICOP1999.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as xw:
    leia.to_excel(xw, sheet_name="Leia-me", index=False); tab.to_excel(xw, sheet_name="Correcoes", index=False); imp.to_excel(xw, sheet_name="Impacto_por_classe", index=False)
wb = openpyxl.load_workbook(out); H = PatternFill("solid", fgColor="1F3A5F")
for ws in wb:
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = H; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for k in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(k)].width = min(max(10, max(len(str(c.value or "")) for c in ws[get_column_letter(k)][:200]) + 2), 70)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
wb["Leia-me"].column_dimensions["A"].width = 140; wb.save(out)
# sincroniza os exemplos do classificador (coicop_ml) com o de-para corrigido
g = pd.read_csv(R / "../coicop_ml/gabarito_paises.csv", dtype=str); nc = novo.coicop
g.loc[g.pais == "Brasil", "coicop"] = g.loc[g.pais == "Brasil", "codigo"].map(nc).fillna(g.coicop)
g.to_csv(R / "../coicop_ml/gabarito_paises.csv", index=False)
b = pd.read_csv(R / "../coicop_ml/brasil_subitens_rotulados.csv", dtype=str); b["coicop"] = b.codigo.map(nc).fillna(b.coicop); b["coicop_desc"] = b.coicop.map(tit).fillna(b.coicop_desc)
b.to_csv(R / "../coicop_ml/brasil_subitens_rotulados.csv", index=False)
print(len(tab), "subitens alterados | peso", round(tab["Peso ago/2026 (%)"].sum(), 2), "% |", novo.coicop.nunique(), "classes de 1999 no Brasil (antes:", m.coicop.nunique(), ")")
print(imp.head(12).to_string(index=False))
