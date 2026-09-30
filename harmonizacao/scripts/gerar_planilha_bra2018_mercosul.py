"""Gera dados/COICOP2018_SNIPC_Brasil_e_Mercosul.xlsx: Brasil (COICOP 2018 x SNIPC) + varredura do Mercosul."""
import re, openpyxl, pandas as pd
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
R = Path(__file__).resolve().parents[1]
def preparar():
    """Subitens do Brasil com peso (último mês do curado), classe 1999, itens do seu rascunho e marca 'alvo' (o que o Claude classifica)."""
    import sys; sys.path.insert(0, str(R/"src")); from harmonizacao import armazem as A
    m = pd.read_csv(A.MAPS/"BRA.csv", dtype=str); m["item"] = m.cod_local.str[:4]
    it = A.ler("item_nacional").set_index("cod_local").nome
    w = A.ler("peso_nacional"); w = w[(w.pais == "BRA") & (w.data == w.data.max())].set_index("cod_local").peso
    m["nome"] = m.cod_local.map(it); m["item_nome"] = m["item"].map(it); m["peso"] = m.cod_local.map(w)
    wb = openpyxl.load_workbook(R/"fontes/brasil/Brasil_rascunho_COICOP2018.xlsx", read_only=True); draft = {}
    for r in wb["Planilha1"].iter_rows(values_only=True):
        cop = str(r[5]).strip() if r[5] is not None else ""
        if not re.fullmatch(r"\d{2}\.\d\.\d", cop): continue             # só linhas com classe (3 níveis)
        for j in range(7, len(r) - 1, 2):
            if r[j] is not None and r[j + 1]:
                c = str(r[j]); draft["1101" if c.startswith("1.1.0") else c] = cop
    m["rasc"] = [draft.get(c) or draft.get(i) for c, i in zip(m.cod_local, m["item"])]
    # itens que o rascunho deixou pendentes/marcou para dividir + itens cujos subitens caem em >1 classe de 1999
    manual = {"2104","6301","5101","5102","5104","7101","7201","8101","8102","8103","8104","1110","1114","2101","3202","3102","3301","1115","4401","4201","2103"}
    multi = set(m.groupby("item").coicop.nunique().loc[lambda s: s > 1].index)
    m["alvo"] = m["item"].isin(manual | multi) | m.rasc.isna()
    return m
m = preparar()
sug = pd.read_csv(R/"mapeamentos/BRA_2018_sugestao.csv", dtype=str).set_index("cod_local")
tb = pd.read_csv(R/"tabelas/coicop_correspondencia_2018_1999.csv", dtype=str)
cls = tb[tb.coicop2018.str.count(r"\.") == 2].drop_duplicates("coicop2018")[["coicop2018", "titulo2018"]]
cls = cls[~cls.coicop2018.str.match(r"1[45]\.")]
cls["titulo2018"] = cls.titulo2018.str.replace(r"\s*\((ND|SD|D|S)\)$", "", regex=True)
tit = dict(zip(cls.coicop2018, cls.titulo2018))
# --- Brasil por subitem
def origem(r):
    if r.cod_local in sug.index: return ("Sugestão Claude" if r.alvo else "Exceção Claude ao seu rascunho")
    return "Seu rascunho (por item)"
m["origem"] = m.apply(origem, axis=1)
m["classe2018"] = [sug.classe2018.get(c) if c in sug.index else r for c, r in zip(m.cod_local, m.rasc)]
m["confianca"] = [sug.confianca.get(c, "—") if c in sug.index else "sua decisão" for c in m.cod_local]
m["segunda"] = [sug.segunda_classe.get(c, "") if c in sug.index else "" for c in m.cod_local]
m["just"] = [sug.justificativa.get(c, "") if c in sug.index else "Item colocado por você na classe" for c in m.cod_local]
m["tit2018"] = m.classe2018.map(tit)
bad = m[m.classe2018.isna()]; assert bad.empty, bad[["cod_local", "nome"]]
m = m.sort_values("peso", ascending=False)
sub = pd.DataFrame({"Código SNIPC": m.cod_local, "Item SNIPC": m.item_nome, "Subitem SNIPC": m.nome, "Peso ago/2026 (%)": m.peso.round(4),
    "COICOP 2018 (classe)": m.classe2018, "Título COICOP 2018": m.tit2018, "Origem": m.origem, "Confiança": m.confianca,
    "Se divide: 2ª classe": m.segunda, "Justificativa": m.just, "COICOP 1999 (seu pré-gabarito)": m.coicop,
    "Sua revisão (OK / classe correta)": ""})
# --- lado a lado: uma linha por classe 2018
lado = []
for c, t in cls.values:
    x = m[m.classe2018 == c].sort_values("peso", ascending=False)
    lado.append({"COICOP 2018": c, "Título": t, "Peso Brasil (%)": round(x.peso.sum(), 4), "Nº subitens": len(x),
                 "Situação": "preenchida" if len(x) else "vazia (sem preço no IPCA)",
                 "Subitens SNIPC (código nome – peso %)": "\n".join(f"{a} {b} – {p:.3f}" for a, b, p in zip(x.cod_local, x.nome, x.peso))})
lado = pd.DataFrame(lado)
# --- revisão por peso (só o que eu classifiquei)
rev = sub[sub["Origem"] != "Seu rascunho (por item)"].copy(); rev["Peso acumulado (%)"] = rev["Peso ago/2026 (%)"].cumsum().round(2)
rev["% do peso a revisar"] = (rev["Peso acumulado (%)"] / rev["Peso ago/2026 (%)"].sum() * 100).round(1)
rev = rev[["Código SNIPC", "Subitem SNIPC", "Peso ago/2026 (%)", "Peso acumulado (%)", "% do peso a revisar", "COICOP 2018 (classe)", "Título COICOP 2018", "Confiança", "Se divide: 2ª classe", "Justificativa", "Sua revisão (OK / classe correta)"]]
# --- Mercosul
F = "Fonte/observação"
merc = pd.DataFrame([
 ["Brasil","IBGE","IPCA (SNIPC), pesos da POF","mensal, 457 categorias até subitem","SNIPC (próprio); COICOP: sem correspondência oficial localizada","Sim: SIDRA 7060 (variação e peso mensal por subitem)","Cesta completa: obtida e conectada ao pipeline","https://sidra.ibge.gov.br/tabela/7060"],
 ["Argentina","INDEC","IPC base dez/2016=100 (nacional + 6 regiões)","pesos até classe COICOP (3 níveis) por região; 59 classes","COICOP 1999 (CCIF), adotada no IPC nacional desde jun/2017","Parcial: pesos por divisão/grupo/classe (Cuadro 7 da metodologia); lista de artigos/produtos não localizada","Extraído: 108 linhas (12 divisões, 37 grupos, 59 classes) × 6 regiões. Pesos são regionais; ponderador nacional não consta nesse quadro","https://www.indec.gob.ar/ftp/cuadros/economia/metodologia_ipc_nacional_2019.pdf"],
 ["Paraguai","BCP","IPC base dez/2017","465 produtos (358 bens, 107 serviços) em 155 subgrupos, 56 grupos, 12 divisões","CCIF/COICOP 1999 (texto da metodologia); códigos próprios de 8 dígitos","Sim: Anexo 2 da metodologia traz cesta completa com ponderações","Extraído e conferido: 465 artigos, 155 subgrupos, 56 grupos, 12 divisões; somas por nível ≈ 100","https://informacionpublica.paraguay.gov.py/public/1970350-Metodologa_IPC_Base_Diciembre_2017_ma2pdf-Metodologa_IPC_Base_Diciembre_2017_ma2.pdf"],
 ["Uruguai","INE","IPC base out/2022=100 (3 cestas: país, Montevidéu, interior)","estrutura publicada em divisão e grupo (segundo a busca); pesos mensais variáveis","CCIF/COICOP 1999","NÃO VERIFICADO: servidores do INE (www5/www7.ine.gub.uy) derrubam a conexão deste ambiente","Sem arquivo obtido. Pendente: tentar do seu computador","https://www.gub.uy/instituto-nacional-estadistica/comunicacion/publicaciones/metodologia-indice-precios-del-consumo-ipc-base-octubre-2022100"],
 ["Bolívia","INE","IPC base 2016 (9 cidades/conurbações)","pesos por divisão e cidade; 12 divisões (alimentação fora do lar como divisão própria)","CCIF (estrutura de 12 divisões adaptada)","Só divisão: PDF de ponderações obtido; RAR com detalhe veio com arquivo vazio","Extraído: 12 divisões × 10 colunas (9 cidades + Bolívia)","https://www.ine.gob.bo/index.php/publicaciones/ponderaciones-ipc-2016/"],
 ["Venezuela","BCV","INPC base dez/2007","362 produtos (segundo a busca)","CCIF","NÃO VERIFICADO: não abri documento primário; continuidade da publicação não checada","Sem arquivo obtido","https://www.bcv.org.ve/"]],
 columns=["País","Instituto","Índice / base","Detalhe da cesta e pesos","Classificador","Pesos por item públicos?","O que foi obtido aqui",F])
# --- Paraguai
L=[l.strip() for l in open(R/"fontes/mercosul/PRY_metodologia_base2017.txt").read().split("\n")]
i0=[i for i,l in enumerate(L) if l.startswith("Anexo 2 – Canasta")][-1]; i1=[i for i,l in enumerate(L) if l.startswith("Anexo 3 – Serie")][-1]
rows=[];i=i0
while i<i1:
    if re.fullmatch(r"0|\d{8,9}",L[i]):                     # códigos do BCP: 0 (total) ou 8-9 dígitos
        j=i+1; d=[]
        while j<i1 and not (re.fullmatch(r"[0-4]",L[j]) and j+1<i1 and re.fullmatch(r"\d+,\d+",L[j+1])): d.append(L[j]); j+=1
        if j<i1 and d: rows.append((L[i]," ".join(d).strip(),int(L[j]),float(L[j+1].replace(",",".")))); i=j+2; continue
    i+=1
pry=pd.DataFrame(rows,columns=["Código BCP","Descrição","Nível (0 total, 1 divisão, 2 grupo, 3 subgrupo, 4 artigo)","Ponderação (%)"])
div=[t for t in pry[pry.iloc[:,2]==1].Descrição]; dmap={c:f"{k+1:02d}" for k,c in enumerate(pry[pry.iloc[:,2]==1]["Código BCP"])}
pry["Divisão COICOP (1 a 12, pela ordem do BCP)"]=[("" if c=="0" else f"{int(c[:-7]):02d}") for c in pry["Código BCP"]]
soma={n:round(pry[pry.iloc[:,2]==n]["Ponderação (%)"].sum(),3) for n in (1,2,3,4)}
# --- Argentina / Bolívia
arg=pd.read_csv(R/"fontes/mercosul/ARG_ponderadores_regionais_2016.csv",dtype={"cod":str}).rename(columns={"cod":"COICOP 1999","desc":"Descrição","nivel":"Nível"})
bol=pd.DataFrame({"Divisão":["Alimentos e bebidas não alcoólicas","Bebidas alcoólicas e tabaco","Vestuário e calçados","Moradia e serviços básicos","Móveis e serviços domésticos","Saúde","Transporte","Comunicações","Recreação e cultura","Educação","Alimentos e bebidas fora do lar","Bens e serviços diversos","TOTAL"],
 "Bolívia (%)":[27.06,0.88,7.56,8.56,6.08,3.55,9.07,5.43,6.22,4.07,13.95,7.55,100.00]})
leia=pd.DataFrame({"Leia-me":[
 "Planilha gerada em 30/09/2026. Brasil: COICOP 2018 × SNIPC (subitem). Mercosul: varredura inicial das cestas publicadas.",
 "",
 "ABA Brasil_por_subitem: os 377 subitens do IPCA com peso de ago/2026. A coluna Origem diz quem decidiu: 'Seu rascunho (por item)' = a classe que você colocou no item; 'Sugestão Claude' = subitens de itens que o seu rascunho deixou pendentes ou se dividem; 'Exceção Claude ao seu rascunho' = 12 subitens em que o texto oficial da COICOP 2018 contradiz a colocação do item inteiro.",
 "ABA Brasil_lado_a_lado: uma linha por classe COICOP 2018 (146 classes, excluídas as divisões 14 e 15), com os subitens SNIPC que caíram nela. Classes vazias = sem preço no IPCA.",
 "ABA Revisao_por_peso: só o que o Claude classificou, do maior para o menor peso, com o peso acumulado. Preencha a última coluna.",
 "Confiança: alta = texto oficial da ONU cita o produto ou o caso é óbvio; média = plausível, com alternativa; baixa = subitem misto ou sem classe clara. 'Se divide' = 2ª classe para repartir o peso quando o subitem cobre duas classes (a proporção ainda não foi definida).",
 "",
 "LIMITES: (1) As sugestões são minhas e NÃO foram revisadas; o objetivo é medir quanto você aceita sem mudar. (2) Classificação só no nível de classe (3 níveis); subclasses (4 níveis) não foram atribuídas. (3) O peso é o de ago/2026; o de-para não depende dele, mas a ordem de revisão sim.",
 "(4) Para 12 subitens o pré-gabarito de 1999 e o rascunho colocam o item todo numa classe, mas o texto de 2018 manda outra: feijão (01.1.7), banana-da-terra (01.1.7), manteiga (01.1.5), bacalhau (01.1.3), leite de coco (01.1.4), alho (01.1.7), entre outros. Vale corrigir o de-para de 1999 também.",
 "(5) Mercosul: Uruguai e Venezuela NÃO foram verificados (ver aba Mercosul_varredura). Paraguai: cesta extraída do PDF e conferida (465 artigos; soma por nível: "+str({k:round(float(v),2) for k,v in soma.items()})+"). O Paraguai ainda NÃO foi classificado em COICOP 2018; só a divisão foi atribuída."]})
out=R/"dados"/"COICOP2018_SNIPC_Brasil_e_Mercosul.xlsx"; out.parent.mkdir(exist_ok=True)
with pd.ExcelWriter(out,engine="openpyxl") as xw:
    for n,d in [("Leia-me",leia),("Brasil_por_subitem",sub),("Brasil_lado_a_lado",lado),("Revisao_por_peso",rev),("Mercosul_varredura",merc),("PRY_cesta",pry),("ARG_pesos_COICOP99",arg),("BOL_pesos_divisao",bol)]: d.to_excel(xw,sheet_name=n,index=False)
wb=openpyxl.load_workbook(out); H=PatternFill("solid",fgColor="1F3A5F"); COR={"alta":"E2F0D9","média":"FFF2CC","baixa":"F8CBAD"}
for ws in wb:
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=H; c.alignment=Alignment(wrap_text=True,vertical="center")
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for k in range(1,ws.max_column+1):
        w=max(len(str(c.value or "")) for c in ws[get_column_letter(k)][:200]); ws.column_dimensions[get_column_letter(k)].width=min(max(10,w+2),60)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
wb["Leia-me"].column_dimensions["A"].width=140
for n in ("Brasil_por_subitem","Revisao_por_peso"):
    ws=wb[n]; hdr=[c.value for c in ws[1]]; k=hdr.index("Confiança")+1
    for row in ws.iter_rows(min_row=2):
        v=row[k-1].value
        if v in COR: row[k-1].fill=PatternFill("solid",fgColor=COR[v])
wb["Brasil_lado_a_lado"].column_dimensions["F"].width=70
wb.save(out)
print(out, "| subitens",len(sub),"| classes preenchidas",(lado["Nº subitens"]>0).sum(),"de",len(lado),"| peso coberto",round(lado["Peso Brasil (%)"].sum(),2))
print("a revisar:",len(rev),"subitens, peso",round(rev["Peso ago/2026 (%)"].sum(),2),"| 80% do peso em",int((rev["% do peso a revisar"]<=80).sum())+1,"subitens")
print("Paraguai somas por nível:",soma)
print(sub.groupby("Origem")["Peso ago/2026 (%)"].agg(["count","sum"]).round(1).to_string())
